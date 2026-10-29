"""GradScout workbook: a permanent, formatted Excel database.

Sheets: Opportunities (shortlist), Applications (what you are applying to), Reviewed (not shortlisted),
Parked (jobs & internships), Search log. Rows are plain dicts keyed by column header; the whole
workbook is rewritten on every save, so formatting is always consistent. Nothing is ever deleted.
"""
import os
import re
import shutil
from collections import Counter
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

import config as cfg

# ---------------------------------------------------------------- opportunity sheets
OPP_COLUMNS = ["Fit", "Title", "Organization", "Category", "Deadline", "Days left", "Open?", "Details", "Eligibility",
               "Fit reason", "Link", "Status", "Notes", "Source", "Added"]
OPP_WIDTHS = [9, 44, 28, 12, 13, 10, 9, 60, 38, 40, 36, 12, 30, 20, 12]
OPP_CENTERED = {"Fit", "Deadline", "Days left", "Open?", "Status", "Added"}
CATEGORIES = ["RA/TA", "Fee waiver", "Fellowship", "Scholarship", "Other"]
STATUSES = ["New", "Interested", "Applying", "Applied", "Rejected", "Skip"]
FITS = ["High", "Medium", "Low"]

# ---------------------------------------------------------------- applications sheet
STAGES = ["Planning", "Preparing", "Submitted", "Interview", "Offer", "Rejected", "Withdrawn"]
REQ_STATES = ["Required", "In progress", "Done", "N/A"]
LEVELS = ["PhD", "Master's", "Other"]
REPLIES = ["Pending", "Positive", "Negative", "No reply"]
DECISIONS = ["Pending", "Funded offer", "Offer", "Waitlisted", "Rejected", "Withdrawn"]
APP_GROUPS = [
    ("Opportunity", "1F4E78", ["Opportunity", "Type", "Provider", "Level", "Link", "Contact"]),
    ("Dates", "2E75B6", ["Opens", "Deadline", "Days left"]),
    ("Progress", "548235", ["Stage", "Next action", "Next action date"]),
    ("Application fee", "BF8F00", ["Fee status", "Fee amount", "Waiver route"]),
    ("Requirements", "7030A0", cfg.REQUIREMENTS[:6] + ["Recommenders"] + cfg.REQUIREMENTS[6:]),
    ("Outcome", "C55A11", ["Submitted on", "Follow-up sent", "Reply", "Interview date", "Decision", "Notes"]),
]
APP_COLUMNS = [c for _, _, cols in APP_GROUPS for c in cols]
APP_DATES = {"Opens", "Deadline", "Next action date", "Submitted on", "Interview date"}
APP_WIDE = {"Opportunity": 40, "Provider": 26, "Link": 30, "Contact": 26, "Next action": 32, "Waiver route": 24,
            "Recommenders": 24, "Notes": 34, "Type": 20}
APP_CHOICES = {"Type": cfg.APP_TYPES, "Level": LEVELS, "Stage": STAGES, "Fee status": cfg.FEE_STATUSES,
               "Follow-up sent": ["Yes", "No"], "Reply": REPLIES, "Decision": DECISIONS,
               **{r: REQ_STATES for r in cfg.REQUIREMENTS}}

ACTIVE, APPS = "Opportunities", "Applications"
REVIEWED, PARKED, LOGSHEET = "Reviewed (not shortlisted)", "Parked (jobs & internships)", "Search log"
FORMULA_HEADERS = {"Days left", "Open?"}
DATE_FORMATS = ("%d %B %Y", "%B %d %Y", "%d %b %Y", "%b %d %Y", "%Y-%m-%d", "%m/%d/%Y")
UNSCORED_REASON = "Unscored: shown but not added"
LAST_ROW = 1000
BACKUPS_KEPT = 30


# ---------------------------------------------------------------- small helpers

def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", str(value or "")).replace(",", " ").replace(".", "")
    text = re.sub(r"\s+", " ", text).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def iso(value) -> str:
    if isinstance(value, (datetime, date)):
        return parse_date(value).isoformat()
    return "" if value is None else str(value)


def norm_link(link: str) -> str:
    return str(link or "").strip().lower().rstrip("/")


def pkey(link: str, title: str) -> str:
    return norm_link(link) or title.strip().lower()


def blank(row: dict, columns: list[str] = OPP_COLUMNS) -> dict:
    return {c: row.get(c, "") for c in columns}


def fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def find(rows: list[dict], title: str, key: str = "Title") -> tuple[dict | None, str]:
    """Find one row by exact title, then title substring, then organization/provider substring."""
    t = title.lower()
    hits = ([r for r in rows if r[key].lower() == t] or [r for r in rows if t in r[key].lower()]
            or [r for r in rows if t in (r.get("Organization") or r.get("Provider") or "").lower()])
    if not hits:
        return None, f"No row matching '{title}'"
    if len(hits) > 1:
        return None, f"'{title}' matches {len(hits)} rows: " + "; ".join(r[key] for r in hits)
    return hits[0], ""


def within(row: dict, rows: list[dict]) -> bool:
    return any(row is r for r in rows)


def opp_sort_key(row: dict):
    fit = FITS.index(row["Fit"]) if row["Fit"] in FITS else len(FITS)
    return fit, parse_date(row["Deadline"]) or date.max, row["Title"].lower()


def app_sort_key(row: dict):
    stage = STAGES.index(row["Stage"]) if row["Stage"] in STAGES else 0
    return stage, parse_date(row["Deadline"]) or date.max, row["Opportunity"].lower()


def all_opportunities(book: dict) -> list[dict]:
    return book["active"] + book["reviewed"] + book["parked"]


# ---------------------------------------------------------------- load / save

def _read_sheet(ws, columns: list[str], header_row: int, key: str) -> list[dict]:
    headers = [c.value for c in ws[header_row]]
    rows = []
    for cells in ws.iter_rows(min_row=header_row + 1):
        row = {}
        for h, c in zip(headers, cells):
            if not h or h in FORMULA_HEADERS:
                continue
            value = c.hyperlink.target if (h == "Link" and c.hyperlink and c.hyperlink.target) else c.value
            row[h] = iso(value)
        if row.get(key):
            rows.append(blank(row, columns))
    return rows


def load() -> dict:
    """Return the workbook as {'active', 'reviewed', 'parked', 'applications'}: lists of row dicts."""
    book = {"active": [], "reviewed": [], "parked": [], "applications": []}
    path = cfg.tracker_path()
    if not path.exists():
        return book
    wb = load_workbook(path)
    first = wb[ACTIVE] if ACTIVE in wb.sheetnames else wb.active
    book["active"] = _read_sheet(first, OPP_COLUMNS, 1, "Title")
    for key, name in (("reviewed", REVIEWED), ("parked", PARKED)):
        if name in wb.sheetnames:
            book[key] = _read_sheet(wb[name], OPP_COLUMNS, 1, "Title")
    if APPS in wb.sheetnames:
        book["applications"] = _read_sheet(wb[APPS], APP_COLUMNS, 2, "Opportunity")
    return book


def _backup(path) -> None:
    """Keep one copy of the workbook per day (the first save of the day), newest BACKUPS_KEPT only."""
    if not path.exists():
        return
    folder = cfg.backups_dir()
    target = folder / f"GradScout-{date.today().isoformat()}.xlsx"
    if target.exists():
        return
    folder.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, target)
    for old in sorted(folder.glob("GradScout-*.xlsx"))[:-BACKUPS_KEPT]:
        old.unlink()


def save(book: dict) -> None:
    path = cfg.tracker_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    _backup(path)
    wb = Workbook()
    ws = wb.active
    ws.title = ACTIVE
    ws.sheet_properties.tabColor = "1F4E78"
    _write_opps(ws, sorted(book["active"], key=opp_sort_key))
    apps = wb.create_sheet(APPS)
    apps.sheet_properties.tabColor = "548235"
    _write_apps(apps, sorted(book["applications"], key=app_sort_key))
    for name, key, color in ((REVIEWED, "reviewed", "7F7F7F"), (PARKED, "parked", "BFBFBF")):
        sheet = wb.create_sheet(name)
        sheet.sheet_properties.tabColor = color
        _write_opps(sheet, sorted(book[key], key=opp_sort_key))
    _write_log(wb.create_sheet(LOGSHEET))
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise RuntimeError(f"Close {path.name} in Excel and retry.")


def refresh() -> None:
    """Rewrite the workbook (e.g. to update the Search log sheet); silently skipped if Excel has it open."""
    try:
        if cfg.tracker_path().exists():
            save(load())
    except RuntimeError:
        pass


# ---------------------------------------------------------------- writers

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header_cell(cell, color: str, dark_text: bool = False) -> None:
    cell.font = Font(bold=True, color="1F1F1F" if dark_text else "FFFFFF")
    cell.fill = fill(color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def _body_cell(cell, centered: bool, formula: bool = False) -> None:
    if isinstance(cell.value, str) and cell.value.startswith("=") and not formula:
        cell.data_type = "s"  # never let scraped text become a formula
    cell.border = BORDER
    cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if centered else "left")


def _link_cell(cell, url: str) -> None:
    if url.startswith("http"):
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")
        cell.alignment = Alignment(vertical="top", wrap_text=False)


def _dropdowns(ws, choices: dict, columns: list[str], first_row: int) -> None:
    for name, options in choices.items():
        letter = get_column_letter(columns.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}{first_row}:{letter}{LAST_ROW}")


def _equals(ws, ref: str, styles: dict) -> None:
    """Conditional colour rules: {value: (fill or None, font colour)}."""
    for value, (bg, fg) in styles.items():
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=[f'"{value}"'],
                                                      fill=fill(bg) if bg else None, font=Font(color=fg, bold=True)))


def _page(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _deadline_rules(ws, date_col: str, days_col: str, first_row: int, cols: str) -> None:
    ws.conditional_formatting.add(cols, FormulaRule(
        formula=[f"AND(ISNUMBER(${days_col}{first_row}),${days_col}{first_row}<0)"],
        fill=fill("E7E6E6"), font=Font(color="7F7F7F")))
    ws.conditional_formatting.add(cols, FormulaRule(
        formula=[f"AND(ISNUMBER(${days_col}{first_row}),${days_col}{first_row}>=0,"
                 f"${days_col}{first_row}<={cfg.load_settings()['alert_days']})"],
        fill=fill("FCE4D6"), font=Font(color="C65911", bold=True)))


def _write_opps(ws, rows: list[dict]) -> None:
    ws.append(OPP_COLUMNS)
    for i, name in enumerate(OPP_COLUMNS, start=1):
        _header_cell(ws.cell(row=1, column=i), "1F4E78")
        ws.column_dimensions[get_column_letter(i)].width = OPP_WIDTHS[i - 1]
    ws.row_dimensions[1].height = 26
    for n, r in enumerate(rows, start=2):
        ws.append([
            r["Fit"], r["Title"], r["Organization"], r["Category"], parse_date(r["Deadline"]) or r["Deadline"],
            f'=IF(ISNUMBER(E{n}),E{n}-TODAY(),"")',
            f'=IF(ISNUMBER(E{n}),IF(E{n}>=TODAY(),"Open","Closed"),"Check")',
            r["Details"], r["Eligibility"], r["Fit reason"], r["Link"], r["Status"], r["Notes"], r["Source"],
            parse_date(r["Added"]) or r["Added"],
        ])
        for name, c in zip(OPP_COLUMNS, ws[n]):
            _body_cell(c, name in OPP_CENTERED, name in FORMULA_HEADERS)
        ws.cell(row=n, column=1).font = Font(bold=True)
        ws.cell(row=n, column=5).number_format = "dd mmm yyyy"
        ws.cell(row=n, column=6).number_format = "0"
        ws.cell(row=n, column=15).number_format = "dd mmm yyyy"
        _link_cell(ws.cell(row=n, column=11), r["Link"])
    _equals(ws, f"A2:A{LAST_ROW}", {"High": ("C6EFCE", "006100"), "Medium": ("FFEB9C", "9C5700"), "Low": ("FFC7CE", "9C0006")})
    status = get_column_letter(OPP_COLUMNS.index("Status") + 1)
    _equals(ws, f"{status}2:{status}{LAST_ROW}", {
        "New": ("DDEBF7", "1F4E78"), "Interested": ("E4DFEC", "5B2C83"), "Applying": ("FCE4D6", "C65911"),
        "Applied": ("C6EFCE", "006100"), "Rejected": ("FFC7CE", "9C0006"), "Skip": ("E7E6E6", "7F7F7F")})
    _deadline_rules(ws, "E", "F", 2, f"E2:F{LAST_ROW}")
    _equals(ws, f"G2:G{LAST_ROW}", {"Open": (None, "006100"), "Closed": ("E7E6E6", "7F7F7F"), "Check": ("FFEB9C", "9C5700")})
    _dropdowns(ws, {"Fit": FITS, "Category": CATEGORIES, "Status": STATUSES}, OPP_COLUMNS, 2)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OPP_COLUMNS))}{max(len(rows) + 1, 2)}"
    _page(ws)


def _write_apps(ws, rows: list[dict]) -> None:
    col = 1
    for group, color, cols in APP_GROUPS:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        ws.cell(row=1, column=col, value=group)
        for i in range(col, col + len(cols)):
            _header_cell(ws.cell(row=1, column=i), color)
        col += len(cols)
    for i, name in enumerate(APP_COLUMNS, start=1):
        cell = ws.cell(row=2, column=i, value=name)
        _header_cell(cell, "D9E1F2", dark_text=True)
        ws.column_dimensions[get_column_letter(i)].width = APP_WIDE.get(name, 14)
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 34
    letter = {name: get_column_letter(i) for i, name in enumerate(APP_COLUMNS, start=1)}
    dl, days = letter["Deadline"], letter["Days left"]
    for n, r in enumerate(rows, start=3):
        values = []
        for name in APP_COLUMNS:
            if name == "Days left":
                values.append(f'=IF(ISNUMBER({dl}{n}),{dl}{n}-TODAY(),"")')
            elif name in APP_DATES:
                values.append(parse_date(r[name]) or r[name])
            else:
                values.append(r[name])
        ws.append(values)
        for name, c in zip(APP_COLUMNS, ws[n]):
            _body_cell(c, name in APP_CHOICES or name in APP_DATES or name in ("Days left", "Fee amount"),
                       name == "Days left")
            if name in APP_DATES:
                c.number_format = "dd mmm yyyy"
        ws[f"{days}{n}"].number_format = "0"
        ws[f"{letter['Opportunity']}{n}"].font = Font(bold=True)
        _link_cell(ws[f"{letter['Link']}{n}"], r["Link"])
    _equals(ws, f"{letter['Stage']}3:{letter['Stage']}{LAST_ROW}", {
        "Planning": ("DDEBF7", "1F4E78"), "Preparing": ("FFEB9C", "9C5700"), "Submitted": ("BDD7EE", "1F4E78"),
        "Interview": ("BDD7EE", "1F4E78"), "Offer": ("C6EFCE", "006100"), "Rejected": ("E7E6E6", "7F7F7F"),
        "Withdrawn": ("E7E6E6", "7F7F7F")})
    first_req, last_req = letter[cfg.REQUIREMENTS[0]], letter[cfg.REQUIREMENTS[-1]]
    _equals(ws, f"{first_req}3:{last_req}{LAST_ROW}", {
        "Done": ("C6EFCE", "006100"), "In progress": ("FFEB9C", "9C5700"), "Required": ("FFC7CE", "9C0006"),
        "N/A": ("E7E6E6", "7F7F7F")})
    _equals(ws, f"{letter['Fee status']}3:{letter['Fee status']}{LAST_ROW}", {
        "Waived": ("C6EFCE", "006100"), "Waiver requested": ("FFEB9C", "9C5700"), "Required": ("FFC7CE", "9C0006"),
        "Paid": ("E7E6E6", "595959"), "Not required": ("E7E6E6", "595959")})
    _equals(ws, f"{letter['Decision']}3:{letter['Decision']}{LAST_ROW}", {
        "Funded offer": ("C6EFCE", "006100"), "Offer": ("C6EFCE", "006100"), "Waitlisted": ("FFEB9C", "9C5700"),
        "Rejected": ("E7E6E6", "7F7F7F"), "Withdrawn": ("E7E6E6", "7F7F7F")})
    _deadline_rules(ws, dl, days, 3, f"{dl}3:{days}{LAST_ROW}")
    na = letter["Next action date"]
    ws.conditional_formatting.add(f"{na}3:{na}{LAST_ROW}", FormulaRule(
        formula=[f"AND(ISNUMBER({na}3),{na}3<TODAY())"], fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(f"{na}3:{na}{LAST_ROW}", FormulaRule(
        formula=[f"AND(ISNUMBER({na}3),{na}3>=TODAY(),{na}3-TODAY()<={cfg.load_settings()['next_action_alert_days']})"],
        fill=fill("FCE4D6"), font=Font(color="C65911", bold=True)))
    _dropdowns(ws, APP_CHOICES, APP_COLUMNS, 3)
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(APP_COLUMNS))}{max(len(rows) + 2, 3)}"
    _page(ws)


def _write_log(ws) -> None:
    headers = ["Source", "Query", "Run at (UTC)", "Window start", "Window end", "Found", "Added"]
    ws.append(headers)
    for i, width in enumerate([18, 30, 22, 22, 22, 8, 8], start=1):
        _header_cell(ws.cell(row=1, column=i), "1F4E78")
        ws.column_dimensions[get_column_letter(i)].width = width
    for run in reversed(cfg.load_log()["runs"]):
        ws.append([run["source"], run["query"], run["run_at"], run["window_start"], run["window_end"],
                   run["found"], "" if run["added"] is None else run["added"]])
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- pending results (safety net)

def stash(results: list[dict]) -> None:
    """Remember search results shown to Claude until they are recorded in the workbook."""
    if results:
        pending = cfg.load_pending()
        for r in results:
            pending[pkey(r.get("link", ""), r["title"])] = r
        cfg.save_pending(pending)


def flush_pending() -> int:
    """Move results that were shown but never recorded into the Reviewed sheet as 'Unscored'."""
    pending = cfg.load_pending()
    if not pending:
        return 0
    book = load()
    everything = all_opportunities(book)
    links = {norm_link(r["Link"]) for r in everything if r["Link"]}
    names = {(r["Title"].lower(), r["Organization"].lower()) for r in everything}
    moved = 0
    for it in pending.values():
        link, org = norm_link(it.get("link", "")), str(it.get("organization") or "")
        if (link and link in links) or (it["title"].lower(), org.lower()) in names:
            continue
        book["reviewed"].append(blank({
            "Title": it["title"], "Organization": org, "Category": it.get("category", "Other"),
            "Deadline": iso(it.get("deadline", "")), "Details": it.get("details", ""), "Fit reason": UNSCORED_REASON,
            "Link": it.get("link", ""), "Status": "New", "Source": it.get("source", ""),
            "Added": date.today().isoformat(),
        }))
        moved += 1
    if moved:
        save(book)
    cfg.save_pending({})
    return moved


# ---------------------------------------------------------------- digest

def digest(book: dict, settings: dict, days: int = 0, today: date | None = None) -> dict:
    """Upcoming deadlines, due next actions and counts, for the home window and the scheduled refresh."""
    today = today or date.today()
    days = days or settings["alert_days"]
    applied = {a["Opportunity"].lower() for a in book["applications"]}
    deadlines, actions = [], []
    for r in book["active"]:
        d = parse_date(r["Deadline"])
        if d and 0 <= (d - today).days <= days and r["Status"] not in ("Skip", "Rejected") \
                and r["Title"].lower() not in applied:
            deadlines.append({"kind": "Opportunity deadline", "title": r["Title"], "fit": r["Fit"],
                              "date": d.isoformat(), "days_left": (d - today).days, "link": r["Link"]})
    for a in book["applications"]:
        if a["Stage"] in ("Rejected", "Withdrawn"):
            continue
        d = parse_date(a["Deadline"])
        if d and 0 <= (d - today).days <= days and a["Stage"] in ("", "Planning", "Preparing"):
            deadlines.append({"kind": "Application deadline", "title": a["Opportunity"], "stage": a["Stage"],
                              "date": d.isoformat(), "days_left": (d - today).days, "link": a["Link"]})
        n = parse_date(a["Next action date"])
        if a["Next action"] and n and (n - today).days <= settings["next_action_alert_days"]:
            actions.append({"title": a["Opportunity"], "next_action": a["Next action"], "date": n.isoformat(),
                            "days_left": (n - today).days, "overdue": n < today})
    fits = Counter(r["Fit"] or "Unscored" for r in book["active"])
    return {
        "as_of": today.isoformat(), "alert_days": days,
        "deadlines_soon": sorted(deadlines, key=lambda x: x["days_left"]),
        "next_actions_due": sorted(actions, key=lambda x: x["days_left"]),
        "applications_by_stage": dict(Counter(a["Stage"] or "Planning" for a in book["applications"])),
        "shortlist_by_fit": {k: fits.get(k, 0) for k in FITS + ["Unscored"]},
        "reviewed": len(book["reviewed"]), "parked": len(book["parked"]),
        "added_last_7_days": [r["Title"] for r in book["active"]
                              if parse_date(r["Added"]) and (today - parse_date(r["Added"])).days <= 7],
    }
