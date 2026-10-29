from datetime import date, timedelta

from openpyxl import load_workbook


def opp(store, title, **kw):
    return store.blank({"Title": title, "Status": "New", **kw})


def test_parse_date_formats(gs):
    p = gs.store.parse_date
    assert p("20 November 2026") == date(2026, 11, 20)
    assert p("November 20, 2026") == date(2026, 11, 20)
    assert p("14th October 2026") == date(2026, 10, 14)
    assert p("10/23/2026") == date(2026, 10, 23)
    assert p("2027-01-15") == date(2027, 1, 15)
    assert p("Rolling") is None


def test_round_trip_and_formatting(gs):
    st = gs.store
    book = {"active": [opp(st, "Alpha PhD", Fit="High", Deadline="2027-01-15", Link="https://a.org/x"),
                       opp(st, "=HYPERLINK(\"evil\")", Fit="Medium")],
            "reviewed": [opp(st, "Low one", Fit="Low")], "parked": [],
            "applications": [st.blank({"Opportunity": "Alpha PhD", "Stage": "Planning", "CV": "Done",
                                       "Deadline": "2027-01-15", "Link": "https://a.org/x"}, st.APP_COLUMNS)]}
    st.save(book)
    loaded = st.load()
    assert [r["Title"] for r in loaded["active"]] == ["Alpha PhD", "=HYPERLINK(\"evil\")"]
    assert loaded["reviewed"][0]["Fit"] == "Low"
    assert loaded["applications"][0]["CV"] == "Done"
    assert loaded["active"][0]["Deadline"] == "2027-01-15"

    wb = load_workbook(gs.cfg.tracker_path())
    assert wb.sheetnames == [st.ACTIVE, st.APPS, st.REVIEWED, st.PARKED, st.LOGSHEET]
    ws = wb[st.ACTIVE]
    assert ws["K2"].hyperlink.target == "https://a.org/x"
    assert ws["G2"].value.startswith("=IF(ISNUMBER(E2)")
    assert ws["B3"].data_type == "s"  # scraped text never becomes a formula
    assert {str(r.sqref) for r in ws.conditional_formatting} >= {"A2:A1000", "E2:F1000", "G2:G1000"}
    assert len(ws.data_validations.dataValidation) == 3
    apps = wb[st.APPS]
    assert [c.value for c in apps[1] if c.value] == [g for g, _, _ in st.APP_GROUPS]
    assert apps.freeze_panes == "B3"


def test_daily_backup_and_pruning(gs):
    st = gs.store
    st.save({"active": [opp(st, "A")], "reviewed": [], "parked": [], "applications": []})
    folder = gs.cfg.backups_dir()
    folder.mkdir(parents=True, exist_ok=True)
    for i in range(35):
        (folder / f"GradScout-2020-01-{i + 1:02d}.xlsx").write_bytes(b"x")
    st.save(st.load())
    backups = sorted(p.name for p in folder.glob("GradScout-*.xlsx"))
    assert len(backups) == st.BACKUPS_KEPT
    assert backups[-1] == f"GradScout-{date.today().isoformat()}.xlsx"


def test_digest(gs):
    st = gs.store
    today = date(2026, 11, 1)
    book = {"active": [opp(st, "Soon", Fit="High", Deadline="2026-11-10"),
                       opp(st, "Far", Fit="High", Deadline="2027-06-01"),
                       opp(st, "Skipped", Fit="Medium", Deadline="2026-11-05", Status="Skip")],
            "reviewed": [], "parked": [],
            "applications": [st.blank({"Opportunity": "App1", "Stage": "Preparing", "Deadline": "2026-11-12",
                                       "Next action": "Email prof", "Next action date": "2026-10-30"}, st.APP_COLUMNS)]}
    d = st.digest(book, gs.cfg.load_settings(), today=today)
    assert [x["title"] for x in d["deadlines_soon"]] == ["Soon", "App1"]
    assert d["next_actions_due"][0]["overdue"] is True
    assert d["shortlist_by_fit"]["High"] == 2
    assert d["applications_by_stage"] == {"Preparing": 1}


def test_flush_pending_moves_unrecorded_to_reviewed(gs):
    st = gs.store
    st.save({"active": [opp(st, "Tracked", Link="https://t.org")], "reviewed": [], "parked": [], "applications": []})
    st.stash([{"title": "Tracked", "link": "https://t.org"}, {"title": "Forgotten", "link": "https://f.org"}])
    assert st.flush_pending() == 1
    reviewed = st.load()["reviewed"]
    assert reviewed[0]["Title"] == "Forgotten" and reviewed[0]["Fit reason"] == st.UNSCORED_REASON
    assert gs.cfg.load_pending() == {}
